import os
import tempfile
import time
import unittest
from pathlib import Path

import openpyxl

from tests.test_phase2c_safe_entrypoint import make_smoke_input


LARGE_ROW_COUNT = 300_000
LARGE_COLUMN_COUNT = 29


def expand_food_workbook(path: Path, row_count: int):
    source = openpyxl.load_workbook(path, read_only=True)
    sheet = source.active
    rows = sheet.iter_rows(values_only=True)
    headers = list(next(rows))
    template = list(next(rows))
    source.close()
    while len(headers) < LARGE_COLUMN_COUNT:
        headers.append(f"synthetic_extra_{len(headers) + 1}")
        template.append("")

    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("Sheet0")
    ws.append(headers)
    for index in range(row_count):
        row = list(template)
        row[2] = f"SYN-LARGE-{index:06d}"
        ws.append(row)
    wb.save(path)


@unittest.skipUnless(os.environ.get("FINANCE_RUN_LARGE_TESTS") == "1", "set FINANCE_RUN_LARGE_TESTS=1 to run large synthetic performance test")
class TestLargeSyntheticPerformance(unittest.TestCase):

    def test_full_pipeline_handles_300k_by_29_synthetic_rows(self):
        from src.pipeline_adapter import AdapterRunRequest, create_run, execute_stage

        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            os.environ["FINANCE_PIPELINE_STATE_DIR"] = str(base / "api-state")
            input_dir = make_smoke_input(base)
            expand_food_workbook(input_dir / "billing_food.xlsx", LARGE_ROW_COUNT)

            started = time.perf_counter()
            run = create_run(AdapterRunRequest("guan", input_dir, base / "out", run_id="guan-large-performance"))
            for stage in ["intake", "normalize", "calculate", "reconcile", "report"]:
                result = execute_stage(run["run_id"], stage)
                self.assertIn(result["status"], {"success", "warning"})
            elapsed = time.perf_counter() - started

            self.assertLess(elapsed, 900)
            summary_path = base / "out" / "runs" / "guan-large-performance" / "result-summary.json"
            self.assertTrue(summary_path.exists())


if __name__ == "__main__":
    unittest.main()
