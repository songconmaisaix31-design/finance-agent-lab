import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from .models import RunSummary


HEADER_FONT = Font(bold=True, size=11)
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT_W = Font(bold=True, size=11, color="FFFFFF")
TOTAL_FONT = Font(bold=True, size=11, color="CC0000")
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)


def _style_header(ws, headers: list, row: int = 1):
    for i, h in enumerate(headers, 1):
        cell = ws.cell(row, i, h)
        cell.font = HEADER_FONT_W
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        cell.border = THIN_BORDER


def _style_cell(ws, row: int, col: int, value=None):
    cell = ws.cell(row, col, value)
    cell.border = THIN_BORDER
    return cell


def _auto_width(ws, min_width: int = 10, max_width: int = 30):
    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        width = min(max(max_len + 4, min_width), max_width)
        ws.column_dimensions[col_letter].width = width


def generate_report(summary: RunSummary, output_path: str):
    wb = openpyxl.Workbook()

    _write_overview(wb, summary)
    _write_income(wb, "02_餐饮收入", summary.food_income_categories, "餐饮")
    _write_income(wb, "03_零售收入", summary.retail_income_categories, "零售")
    _write_expense(wb, "04_餐饮支出", summary.food_delivery_categories, summary.run_info.food_hq_fee, "餐饮")
    _write_expense(wb, "05_零售支出", summary.retail_delivery_categories, summary.run_info.retail_hq_fee, "零售")
    _write_crowd_detail(wb, summary)
    _write_anomalies(wb, summary.anomalies)
    _write_run_info(wb, summary)

    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    wb.save(output_path)


def _write_overview(wb, summary: RunSummary):
    ws = wb.active
    ws.title = "01_总览"
    ri = summary.run_info

    headers = ["项目", "餐饮", "零售"]
    _style_header(ws, headers)

    rows = [
        ("收入合计", ri.food_income_total, ri.retail_income_total),
        ("毛G 1% 总部抽点", ri.food_hq_fee, ri.retail_hq_fee),
        ("配送成本合计", ri.food_delivery_total, ri.retail_delivery_total),
        ("总支出", round(ri.food_hq_fee + ri.food_delivery_total, 2),
         round(ri.retail_hq_fee + ri.retail_delivery_total, 2)),
        ("净收入", ri.food_net, ri.retail_net),
    ]

    for i, (label, fv, rv) in enumerate(rows, 2):
        _style_cell(ws, i, 1, label)
        _style_cell(ws, i, 2, fv)
        _style_cell(ws, i, 3, rv)
        if label == "净收入":
            ws.cell(i, 1).font = TOTAL_FONT
            ws.cell(i, 2).font = TOTAL_FONT
            ws.cell(i, 3).font = TOTAL_FONT

    ws.merge_cells("A1:C1")
    _style_cell(ws, 1, 1, f"永城 2026-06-21 收入与支出核算 - 运行状态: {ri.status}")
    ws.cell(1, 1).font = Font(bold=True, size=14)

    _auto_width(ws)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:C{len(rows) + 1}"


def _write_income(wb, sheet_name: str, categories: list, label: str):
    ws = wb.create_sheet(sheet_name)
    headers = ["收入类型", "收入金额", "去重订单量", "单均收入"]
    _style_header(ws, headers)

    total_amount = 0.0
    total_orders = 0
    for i, c in enumerate(categories, 2):
        _style_cell(ws, i, 1, c.standard_name)
        _style_cell(ws, i, 2, c.total_amount)
        _style_cell(ws, i, 3, c.distinct_order_count)
        _style_cell(ws, i, 4, c.avg_per_order)
        total_amount += c.total_amount
        total_orders += c.distinct_order_count

    tr = len(categories) + 2
    avg_all = round(total_amount / total_orders, 2) if total_orders > 0 else 0.0
    for j, v in enumerate([f"{label}合计", round(total_amount, 2), total_orders, avg_all], 1):
        cell = _style_cell(ws, tr, j, v)
        cell.font = TOTAL_FONT

    _auto_width(ws)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:D{tr}"


def _write_expense(wb, sheet_name: str, delivery_categories: list, hq_fee: float, label: str):
    ws = wb.create_sheet(sheet_name)
    headers = ["支出类型", "类别", "单量", "单均成本", "总支出"]
    _style_header(ws, headers)

    row = 2
    _style_cell(ws, row, 1, "总部抽点")
    _style_cell(ws, row, 2, "毛G × 1%")
    _style_cell(ws, row, 3, "-")
    _style_cell(ws, row, 4, "-")
    _style_cell(ws, row, 5, hq_fee)
    row += 1

    cat_names = {
        "zhuansong_pintuan": "专送拼团",
        "zhuansong_normal": "专送正常单",
        "zhongbao_pintuan": "众包拼团",
        "zhongbao_normal": "众包正常单",
    }

    total_delivery = 0.0
    for c in delivery_categories:
        cn = cat_names.get(c.category, c.category)
        _style_cell(ws, row, 1, cn)
        _style_cell(ws, row, 2, "固定单价" if c.unit_cost > 0 else "众包成本")
        _style_cell(ws, row, 3, c.order_count)
        display_avg = c.unit_cost if c.unit_cost > 0 else (
            round(c.total_cost / c.order_count, 2) if c.order_count > 0 else 0.0
        )
        _style_cell(ws, row, 4, display_avg)
        _style_cell(ws, row, 5, c.total_cost)
        total_delivery += c.total_cost
        row += 1

    total_expense = round(hq_fee + total_delivery, 2)
    _style_cell(ws, row, 1, f"{label}支出合计")
    _style_cell(ws, row, 2, "")
    _style_cell(ws, row, 3, "")
    _style_cell(ws, row, 4, "")
    _style_cell(ws, row, 5, total_expense)
    for j in range(1, 6):
        ws.cell(row, j).font = TOTAL_FONT

    _auto_width(ws)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{row}"


def _write_crowd_detail(wb, summary: RunSummary):
    ws = wb.create_sheet("06_众包关联明细")
    headers = ["业务", "类别", "运力线", "订单量", "总成本", "单均成本"]
    _style_header(ws, headers)

    row = 2
    for label, cats in [("餐饮", summary.food_delivery_categories), ("零售", summary.retail_delivery_categories)]:
        for c in cats:
            if c.crowd_breakdown:
                cat_names = {
                    "zhongbao_pintuan": "众包拼团",
                    "zhongbao_normal": "众包正常单",
                }
                cn = cat_names.get(c.category, c.category)
                _style_cell(ws, row, 1, label)
                _style_cell(ws, row, 2, cn)
                _style_cell(ws, row, 3, f"成功关联={c.matched_count} 未关联={c.unmatched_count} 重复={c.duplicate_count}")
                _style_cell(ws, row, 4, c.order_count)
                _style_cell(ws, row, 5, c.total_cost)
                _style_cell(ws, row, 6, round(c.total_cost / c.order_count, 2) if c.order_count > 0 else 0.0)
                row += 1

                for clb in c.crowd_breakdown:
                    _style_cell(ws, row, 1, "")
                    _style_cell(ws, row, 2, "")
                    _style_cell(ws, row, 3, clb.capacity_line)
                    _style_cell(ws, row, 4, clb.order_count)
                    _style_cell(ws, row, 5, clb.total_cost)
                    _style_cell(ws, row, 6, clb.avg_per_order)
                    row += 1

    _auto_width(ws)
    ws.freeze_panes = "A2"


def _write_anomalies(wb, anomalies: list):
    ws = wb.create_sheet("07_异常数据")
    if not anomalies:
        _style_cell(ws, 1, 1, "无异常数据")
        return

    keys = ["type", "source", "row_number", "category", "order_id", "service_package", "amount", "detail"]
    headers = [k for k in keys if any(k in a for a in anomalies)]
    _style_header(ws, headers)

    for i, a in enumerate(anomalies, 2):
        for j, k in enumerate(headers, 1):
            _style_cell(ws, i, j, a.get(k, ""))

    _auto_width(ws)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(anomalies)+1}"


def _write_run_info(wb, summary: RunSummary):
    ws = wb.create_sheet("08_运行信息")
    ri = summary.run_info

    info_rows = [
        ("运行ID", ri.run_id),
        ("开始时间", ri.start_time),
        ("结束时间", ri.end_time),
        ("运行状态", ri.status),
        ("", ""),
        ("--- 输入文件 ---", ""),
    ]
    for name, finfo in summary.file_info.items():
        info_rows.append((name, f"大小={finfo['size_bytes']} 修改时间={finfo['modified']}"))

    info_rows.append(("", ""))
    info_rows.append(("--- 原始/有效行数 ---", ""))
    for name, counts in summary.raw_row_counts.items():
        info_rows.append((name, f"原始={counts['raw']} 有效={counts['valid']} 异常={counts['anomaly']}"))

    info_rows.append(("", ""))
    info_rows.append(("--- 对账检查 ---", ""))
    for check in summary.reconciliation_checks:
        status = "PASS" if check["passed"] else "FAIL"
        info_rows.append((check["check"], f"{status} expected={check.get('expected','')} actual={check.get('actual','')}"))

    _style_header(ws, ["项目", "值"])
    for i, (k, v) in enumerate(info_rows, 2):
        _style_cell(ws, i, 1, k)
        _style_cell(ws, i, 2, v)

    _auto_width(ws, max_width=80)
    ws.freeze_panes = "A2"
