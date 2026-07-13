import csv
import os
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Optional

import openpyxl


@dataclass
class NormalizedBillingRow:
    source_file: str
    source_sheet: str
    source_row: int
    billing_date: str
    business_category: str
    order_id: str
    order_type: str
    delivery_type: str
    service_package_type: str
    settlement_amount: Optional[Decimal]
    gross_transaction_amount: Optional[Decimal]
    net_transaction_amount: Optional[Decimal]
    merchant_id: str
    merchant_name: str
    completed_at: str
    remark: str
    agent_delivery_subsidy: Optional[Decimal] = None
    amount_valid: bool = True
    amount_error: str = ""


CROWD_COST_FIELDNAMES = [
    "source_file", "source_sheet", "source_row",
    "date", "delivery_status", "is_retail", "city",
    "is_group_order", "completed_orders", "total_cost",
]


def normalize_value(val):
    if val is None:
        return ""
    s = str(val).strip()
    return s


def parse_decimal(val):
    if val is None:
        return None, False, "empty"
    try:
        s = str(val).strip()
        if s == "" or s == "-":
            return None, False, "empty"
        s = s.replace("%", "")
        s = s.replace(",", "")
        return Decimal(s), True, ""
    except (InvalidOperation, ValueError):
        return None, False, f"invalid_numeric: {val}"


def load_billing_rows(
    filepath: str,
    field_map,
    business_category: str,
) -> list[NormalizedBillingRow]:
    cm = field_map.column_map
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb[field_map.source_sheet]

    rows = []
    invalid_amounts = []

    for r_idx, row_cells in enumerate(
        ws.iter_rows(min_row=field_map.data_start_row, values_only=False),
        start=field_map.data_start_row,
    ):
        row_vals = [cell.value for cell in row_cells]

        if all(v is None or str(v).strip() == "" for v in row_vals):
            continue

        order_id = normalize_value(row_vals[cm["订单号"]] if "订单号" in cm and len(row_vals) > cm["订单号"] else None)

        settlement_str = row_vals[cm["结算金额"]] if "结算金额" in cm and len(row_vals) > cm["结算金额"] else None
        settlement, settlement_ok, settlement_err = parse_decimal(settlement_str)

        gross_str = row_vals[cm["毛交易额"]] if "毛交易额" in cm and len(row_vals) > cm["毛交易额"] else None
        gross, gross_ok, gross_err = parse_decimal(gross_str)

        net_str = row_vals[cm["净交易额"]] if "净交易额" in cm and len(row_vals) > cm["净交易额"] else None
        net, net_ok, net_err = parse_decimal(net_str)

        amount_valid = settlement_ok and gross_ok and net_ok
        amount_error = ""
        if not settlement_ok:
            amount_error += f"settlement:{settlement_err}; "
        if not gross_ok:
            amount_error += f"gross:{gross_err}; "
        if not net_ok:
            amount_error += f"net:{net_err}; "

        if not amount_valid:
            invalid_amounts.append({
                "source_file": os.path.basename(filepath),
                "source_row": r_idx,
                "order_id": order_id,
                "settlement_raw": str(settlement_str),
                "gross_raw": str(gross_str),
                "net_raw": str(net_str),
                "error": amount_error.strip(),
            })

        agent_subsidy = None
        if "代理商配送费活动补贴" in cm:
            subsidy_raw = row_vals[cm["代理商配送费活动补贴"]] if len(row_vals) > cm["代理商配送费活动补贴"] else None
            subsidy_val, subsidy_ok, _ = parse_decimal(subsidy_raw)
            agent_subsidy = subsidy_val if subsidy_ok else None

        row = NormalizedBillingRow(
            source_file=os.path.basename(filepath),
            source_sheet=field_map.source_sheet,
            source_row=r_idx,
            billing_date=normalize_value(row_vals[cm["账单日期"]] if "账单日期" in cm and len(row_vals) > cm["账单日期"] else None),
            business_category=business_category,
            order_id=order_id,
            order_type=normalize_value(row_vals[cm["订单类型"]] if "订单类型" in cm and len(row_vals) > cm["订单类型"] else ""),
            delivery_type=normalize_value(row_vals[cm["配送方式"]] if "配送方式" in cm and len(row_vals) > cm["配送方式"] else None),
            service_package_type=normalize_value(row_vals[cm["服务包类型"]] if "服务包类型" in cm and len(row_vals) > cm["服务包类型"] else None),
            settlement_amount=settlement,
            gross_transaction_amount=gross,
            net_transaction_amount=net,
            merchant_id=normalize_value(row_vals[cm["商户ID"]] if "商户ID" in cm and len(row_vals) > cm["商户ID"] else ""),
            merchant_name=normalize_value(row_vals[cm["商户名称"]] if "商户名称" in cm and len(row_vals) > cm["商户名称"] else ""),
            completed_at=normalize_value(row_vals[cm["订单完成时间"]] if "订单完成时间" in cm and len(row_vals) > cm["订单完成时间"] else ""),
            remark=normalize_value(row_vals[cm["备注"]] if "备注" in cm and len(row_vals) > cm["备注"] else ""),
            agent_delivery_subsidy=agent_subsidy,
            amount_valid=amount_valid,
            amount_error=amount_error.strip(),
        )
        rows.append(row)

    wb.close()
    return rows


def write_normalized_csv(rows: list[NormalizedBillingRow], output_path: str):
    fieldnames = [
        "source_file", "source_sheet", "source_row",
        "billing_date", "business_category", "order_id",
        "order_type", "delivery_type", "service_package_type",
        "settlement_amount", "gross_transaction_amount",
        "net_transaction_amount", "merchant_id", "merchant_name",
        "completed_at", "remark", "agent_delivery_subsidy",
    ]
    with open(output_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for r in rows:
            writer.writerow({
                "source_file": r.source_file,
                "source_sheet": r.source_sheet,
                "source_row": r.source_row,
                "billing_date": r.billing_date,
                "business_category": r.business_category,
                "order_id": r.order_id,
                "order_type": r.order_type,
                "delivery_type": r.delivery_type,
                "service_package_type": r.service_package_type,
                "settlement_amount": str(r.settlement_amount) if r.settlement_amount is not None else "",
                "gross_transaction_amount": str(r.gross_transaction_amount) if r.gross_transaction_amount is not None else "",
                "net_transaction_amount": str(r.net_transaction_amount) if r.net_transaction_amount is not None else "",
                "merchant_id": r.merchant_id,
                "merchant_name": r.merchant_name,
                "completed_at": r.completed_at,
                "remark": r.remark,
                "agent_delivery_subsidy": str(r.agent_delivery_subsidy) if r.agent_delivery_subsidy is not None else "",
            })


def load_crowd_cost_rows(
    filepath: str,
    field_map,
) -> list[dict]:
    cm = field_map.column_map
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    ws = wb[field_map.source_sheet]

    col_date = cm["日期"]
    col_status = cm["配送状态"]
    col_retail = cm["是否零售"]
    col_city = cm["合并区县名称"]
    col_group = cm["is_group_order"]
    col_completed = cm["有效完成单"]
    col_cost = cm["总成本"]

    max_col_needed = max(col_date, col_status, col_retail, col_city,
                         col_group, col_completed, col_cost) + 1

    rows = []
    for r_idx, row_vals in enumerate(
        ws.iter_rows(min_row=field_map.data_start_row, max_col=max_col_needed, values_only=True),
        start=field_map.data_start_row,
    ):
        if all(v is None or str(v).strip() == "" for v in row_vals[:5]):
            continue

        date_val = normalize_value(row_vals[col_date] if len(row_vals) > col_date else None)

        is_retail = 0
        if len(row_vals) > col_retail and row_vals[col_retail] is not None:
            try:
                is_retail = int(float(str(row_vals[col_retail])))
            except (ValueError, TypeError):
                is_retail = 0

        is_group = 0
        if len(row_vals) > col_group and row_vals[col_group] is not None:
            try:
                is_group = int(float(str(row_vals[col_group])))
            except (ValueError, TypeError):
                is_group = 0

        completed = Decimal("0")
        if len(row_vals) > col_completed and row_vals[col_completed] is not None:
            try:
                completed = Decimal(str(row_vals[col_completed]))
            except (InvalidOperation, ValueError):
                completed = Decimal("0")

        total_cost = Decimal("0")
        if len(row_vals) > col_cost and row_vals[col_cost] is not None:
            try:
                total_cost = Decimal(str(row_vals[col_cost]))
            except (InvalidOperation, ValueError):
                total_cost = Decimal("0")

        rows.append({
            "source_file": os.path.basename(filepath),
            "source_sheet": field_map.source_sheet,
            "source_row": r_idx,
            "date": date_val,
            "delivery_status": normalize_value(row_vals[col_status] if len(row_vals) > col_status else None),
            "is_retail": is_retail,
            "city": normalize_value(row_vals[col_city] if len(row_vals) > col_city else None),
            "is_group_order": is_group,
            "completed_orders": completed,
            "total_cost": total_cost,
        })

    wb.close()
    return rows


def write_crowd_cost_csv(rows: list[dict], output_path: str):
    with open(output_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=CROWD_COST_FIELDNAMES)
        writer.writeheader()
        for r in rows:
            writer.writerow({
                "source_file": r["source_file"],
                "source_sheet": r["source_sheet"],
                "source_row": r["source_row"],
                "date": r["date"],
                "delivery_status": r["delivery_status"],
                "is_retail": r["is_retail"],
                "city": r["city"],
                "is_group_order": r["is_group_order"],
                "completed_orders": str(r["completed_orders"]),
                "total_cost": str(r["total_cost"]),
            })
