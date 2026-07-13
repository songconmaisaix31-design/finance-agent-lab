import openpyxl
import yaml
from pathlib import Path
from datetime import datetime
from .models import BillingRow, CrowdCostRow


def load_config(path: str) -> dict:
    with open(path, "r", encoding="utf-8") as f:
        return yaml.safe_load(f)


def normalize_str(s) -> str:
    if s is None:
        return ""
    return str(s).strip()


def safe_float(v) -> float:
    if v is None:
        return 0.0
    try:
        return float(v)
    except (ValueError, TypeError):
        return 0.0


def safe_int(v) -> int:
    if v is None:
        return 0
    try:
        return int(float(str(v)))
    except (ValueError, TypeError):
        return 0


def _get_col(col_mapping: dict, name: str) -> int:
    return col_mapping["columns"][name] - 1


def load_food_billing(filepath: str, config: dict) -> list[BillingRow]:
    cm = config["food_billing"]
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb[cm["sheet"]]
    data_start = cm["data_start_row"]

    c_order_id = _get_col(cm, "order_id")
    c_settlement = _get_col(cm, "settlement_amount")
    c_delivery = _get_col(cm, "delivery_method")
    c_svc = _get_col(cm, "service_package")
    c_gross = _get_col(cm, "gross_transaction")
    c_biz = _get_col(cm, "business_type")
    c_otype = _get_col(cm, "order_type")

    rows = []
    for r_idx, row in enumerate(ws.iter_rows(min_row=data_start, values_only=True), start=data_start):
        if not row:
            continue

        order_id = normalize_str(row[c_order_id] if len(row) > c_order_id else None)
        settlement = safe_float(row[c_settlement] if len(row) > c_settlement else None)
        delivery = normalize_str(row[c_delivery] if len(row) > c_delivery else None)
        svc = normalize_str(row[c_svc] if len(row) > c_svc else None)
        gross = safe_float(row[c_gross] if len(row) > c_gross else None)
        biz_type = normalize_str(row[c_biz] if len(row) > c_biz else None)
        order_type = normalize_str(row[c_otype] if len(row) > c_otype else None)

        if not order_id and settlement == 0.0:
            continue

        rows.append(BillingRow(
            source="food",
            row_number=r_idx,
            order_id=order_id,
            settlement_amount=settlement,
            delivery_method=delivery,
            service_package=svc,
            gross_transaction=gross,
            business_type=biz_type,
            order_type=order_type,
        ))

    wb.close()
    return rows


def load_retail_billing(filepath: str, config: dict) -> list[BillingRow]:
    cm = config["retail_billing"]
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb[cm["sheet"]]
    data_start = cm["data_start_row"]

    c_order_id = _get_col(cm, "order_id")
    c_settlement = _get_col(cm, "settlement_amount")
    c_delivery = _get_col(cm, "delivery_method")
    c_svc = _get_col(cm, "service_package")
    c_gross = _get_col(cm, "gross_transaction")
    c_biz = _get_col(cm, "business_type")
    c_otype = _get_col(cm, "order_type")

    rows = []
    for r_idx, row in enumerate(ws.iter_rows(min_row=data_start, values_only=True), start=data_start):
        if not row:
            continue

        order_id = normalize_str(row[c_order_id] if len(row) > c_order_id else None)
        settlement = safe_float(row[c_settlement] if len(row) > c_settlement else None)
        delivery = normalize_str(row[c_delivery] if len(row) > c_delivery else None)
        svc = normalize_str(row[c_svc] if len(row) > c_svc else None)
        gross = safe_float(row[c_gross] if len(row) > c_gross else None)
        biz_type = normalize_str(row[c_biz] if len(row) > c_biz else None)
        order_type = normalize_str(row[c_otype] if len(row) > c_otype else None)

        if not order_id and settlement == 0.0:
            continue

        rows.append(BillingRow(
            source="retail",
            row_number=r_idx,
            order_id=order_id,
            settlement_amount=settlement,
            delivery_method=delivery,
            service_package=svc,
            gross_transaction=gross,
            business_type=biz_type,
            order_type=order_type,
        ))

    wb.close()
    return rows


def load_crowd_cost(filepath: str, config: dict) -> list[CrowdCostRow]:
    cm = config["crowd_cost"]
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb[cm["sheet"]]
    data_start = cm["data_start_row"]

    c_waybill = _get_col(cm, "waybill_id")
    c_order = _get_col(cm, "order_id")
    c_cl = _get_col(cm, "capacity_line")
    c_cost = _get_col(cm, "total_cost")
    c_retail = _get_col(cm, "is_retail")
    c_pintuan = _get_col(cm, "is_pintuan")
    c_status = _get_col(cm, "delivery_status")
    c_date = _get_col(cm, "date")

    rows = []
    for r_idx, row in enumerate(ws.iter_rows(min_row=data_start, values_only=True), start=data_start):
        if not row:
            continue

        waybill_id = normalize_str(row[c_waybill] if len(row) > c_waybill else None)
        order_id = normalize_str(row[c_order] if len(row) > c_order else None)
        capacity_line = normalize_str(row[c_cl] if len(row) > c_cl else None)
        total_cost = safe_float(row[c_cost] if len(row) > c_cost else None)
        is_retail = safe_int(row[c_retail] if len(row) > c_retail else None)
        is_pintuan = safe_int(row[c_pintuan] if len(row) > c_pintuan else None)
        delivery_status = normalize_str(row[c_status] if len(row) > c_status else None)
        date = normalize_str(row[c_date] if len(row) > c_date else None)

        if not waybill_id and not order_id and total_cost == 0.0:
            continue

        rows.append(CrowdCostRow(
            row_number=r_idx,
            waybill_id=waybill_id,
            order_id=order_id,
            capacity_line=capacity_line,
            total_cost=total_cost,
            is_retail=is_retail,
            is_pintuan=is_pintuan,
            delivery_status=delivery_status,
            date=date,
        ))

    wb.close()
    return rows


def get_file_info(filepath: str) -> dict:
    path = Path(filepath)
    stat = path.stat()
    return {
        "name": path.name,
        "path": str(path.resolve()),
        "size_bytes": stat.st_size,
        "modified": datetime.fromtimestamp(stat.st_mtime).isoformat(),
    }
